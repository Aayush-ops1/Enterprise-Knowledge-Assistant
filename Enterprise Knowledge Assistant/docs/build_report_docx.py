"""Build project_report.docx from project_report.html.

Single source of truth is the HTML report; this script converts it into an
editable Word document with real Word styles (Heading 1/2/3, TOC field, page
numbers, shaded code blocks, styled tables).

Usage:
    python docs/build_report_docx.py
Output:
    project_report.docx
"""
from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
HTML_SRC = ROOT / "project_report.html"
DOCX_OUT = ROOT / "project_report.docx"

# ---------------------------------------------------------------- palette
NAVY_0 = "0A0D16"
NAVY_1 = "0E1220"
NAVY_3 = "1B2136"
NAVY_4 = "222A44"
GOLD = "9A6B00"
GOLD_BRIGHT = "E8B23A"
INK = "14181F"
INK_SOFT = "4A5568"
GRAY = "8B93A5"
CODE_BG = "F2F4F8"

PILL_COLORS = {"green": "2F9E63", "blue": "3B6FD4", "gold": "9A6B00",
               "purple": "7A4FD0", "red": "D64545"}
NOTE_BG = {"blue": "E3ECFC", "gold": "FAF0D7", "green": "E2F5EA", "red": "FDE8E8"}
NOTE_BORDER = {"blue": "3B6FD4", "gold": "E8B23A", "green": "2F9E63", "red": "D64545"}
BOX_BG = {"gold": "9A6B00", "blue": "1D3F85", "green": "17663F",
          "purple": "4B2F8F", "red": "8F2D2D", None: NAVY_0}

# strip emojis (decorative) but keep typographic symbols (→ ▼ · ⊕ …)
EMOJI = re.compile(
    "[\U0001F000-\U0001FAFF\u2600-\u27BF\uFE0F\u2B00-\u2BFF\u25C8]"
)  # \u25C8 = decorative diamond in chapter headings
HTML_ENTITY = re.compile(r"&(?:[a-zA-Z]+|#[0-9]+|#x[0-9a-fA-F]+);")


def clean(text: str) -> str:
    """Strip decorative emoji; keep whitespace so inline boundaries survive.
    Whitespace collapsing happens in add_inline()."""
    text = EMOJI.sub("", text)
    return text.replace("\xa0", " ")


# ------------------------------------------------------- minimal HTML tree
class TreeBuilder(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = {"tag": "body", "attrs": {}, "children": []}
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = {"tag": tag, "attrs": dict(attrs), "children": []}
        self.stack[-1]["children"].append(node)
        self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        node = {"tag": tag, "attrs": dict(attrs), "children": []}
        self.stack[-1]["children"].append(node)

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i]["tag"] == tag:
                del self.stack[i:]
                break

    def handle_data(self, data):
        self.stack[-1]["children"].append(clean(data))


# -------------------------------------------------------------- document
doc = Document()
doc.core_properties.title = "EKA — Enterprise Knowledge Assistant — Project Report"
doc.core_properties.author = "[Student Name]"

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2)


def _shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _border_left(p, color, sz=22):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(sz))
    left.set(qn("w:space"), "4")
    left.set(qn("w:color"), color)
    pBdr.append(left)
    pPr.append(pBdr)


def _shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


# -------------------------------------------------------------- styles
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

h1 = doc.styles["Heading 1"]
h1.font.name = "Calibri"
h1.font.size = Pt(16)
h1.font.bold = True
h1.font.color.rgb = RGBColor.from_string(NAVY_1)
h1.paragraph_format.space_before = Pt(0)
h1.paragraph_format.space_after = Pt(10)
h1.paragraph_format.keep_with_next = True

h2 = doc.styles["Heading 2"]
h2.font.name = "Calibri"
h2.font.size = Pt(13)
h2.font.bold = True
h2.font.color.rgb = RGBColor.from_string(NAVY_3)
h2.paragraph_format.space_before = Pt(12)
h2.paragraph_format.space_after = Pt(5)
h2.paragraph_format.keep_with_next = True

h3 = doc.styles["Heading 3"]
h3.font.name = "Calibri"
h3.font.size = Pt(11.5)
h3.font.bold = True
h3.font.color.rgb = RGBColor.from_string(NAVY_4)
h3.paragraph_format.space_before = Pt(8)
h3.paragraph_format.space_after = Pt(3)
h3.paragraph_format.keep_with_next = True

code_style = doc.styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
code_style.base_style = doc.styles["Normal"]
code_style.font.name = "Consolas"
code_style.font.size = Pt(8.5)
code_style.paragraph_format.space_after = Pt(0)
code_style.paragraph_format.line_spacing = 1.05
code_style.paragraph_format.left_indent = Cm(0.3)

bullet = doc.styles["List Bullet"]
bullet.paragraph_format.space_after = Pt(2)
num = doc.styles["List Number"]
num.paragraph_format.space_after = Pt(2)


# ------------------------------------------------------------ low-level
def _fmt(run, bold=False, italic=False, color=None, size=None, mono=False):
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    if mono:
        run.font.name = "Consolas"


WS = re.compile(r"\s+")


def add_inline(par, children, bold=False, italic=False, color=None, size=None, mono=False):
    """Recursively render inline HTML children into one paragraph.

    HTML source is pretty-printed, so text nodes carry formatting whitespace.
    Rules: collapse every whitespace run to one space; drop leading space at
    the paragraph start or right after another text run; keep the space when
    a text run follows an element ("<span>1</span> Introduction")."""
    prev_text = False
    at_start = True
    for ch in children:
        if isinstance(ch, str):
            s = WS.sub(" ", ch)
            if at_start or prev_text:
                s = s.lstrip()
            if s:
                run = par.add_run(s)
                _fmt(run, bold=bold, italic=italic, color=color, size=size, mono=mono)
                prev_text = True
                at_start = False
            continue
        tag, attrs, kids = ch["tag"], ch["attrs"], ch["children"]
        if tag in ("strong", "b"):
            add_inline(par, kids, bold=True, italic=italic, color=color, size=size, mono=mono)
        elif tag in ("em", "i"):
            add_inline(par, kids, bold=bold, italic=True, color=color, size=size, mono=mono)
        elif tag == "code":
            add_inline(par, kids, bold=bold, italic=italic, color=color, size=size, mono=True)
        elif tag == "span":
            cls = attrs.get("class", "")
            span_color = None
            span_bold = bold
            if "pill" in cls:
                for k, v in PILL_COLORS.items():
                    if k in cls:
                        span_color = v
                        span_bold = True
            elif "accent" in cls or "num" in cls:
                span_color = GOLD
            add_inline(par, kids, bold=span_bold, italic=italic,
                       color=span_color or color, size=size, mono=mono)
        elif tag == "a":
            add_inline(par, kids, bold=bold, italic=italic, color=color, size=size, mono=mono)
        elif tag == "br":
            par.add_run().add_break()
        else:
            add_inline(par, kids, bold=bold, italic=italic, color=color, size=size, mono=mono)
        if tag != "br":
            prev_text = False
        at_start = False


def text_of(children) -> str:
    out = []
    for ch in children:
        if isinstance(ch, str):
            out.append(ch)
        elif ch["tag"] == "br":
            out.append("\n")
        else:
            out.append(text_of(ch["children"]))
    return "".join(out)


def page_break():
    doc.add_page_break()


def heading(level, children):
    style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    p = doc.add_paragraph(style=style)
    add_inline(p, children)
    return p


def para(children, align=None, color=None, size=None, italic=False, bold=False,
         space_after=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    add_inline(p, children, color=color, size=size, italic=italic, bold=bold)
    return p


def code_block(children):
    txt = text_of(children).strip("\n")
    for line in txt.split("\n"):
        p = doc.add_paragraph(style="CodeBlock")
        run = p.add_run(line)
        _fmt(run, mono=True, color=INK)
        _shade_paragraph(p, CODE_BG)
    # small breathing room after the block
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    sp_runs = sp.add_run("")
    _fmt(sp_runs, size=2)


def _dicts(children):
    return [c for c in children if isinstance(c, dict)]


def render_table(node):
    rows = [c for c in _dicts(node["children"]) if c["tag"] == "tr"]
    if not rows:
        return
    ncols = max(len([c for c in _dicts(r["children"]) if c["tag"] in ("td", "th")]) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, r in enumerate(rows):
        cells = [c for c in _dicts(r["children"]) if c["tag"] in ("td", "th")]
        for ci, cell in enumerate(cells):
            c = t.cell(ri, ci)
            if ri == 0:
                _shade_cell(c, NAVY_1)
            par = c.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            header = cell["tag"] == "th"
            if header:
                add_inline(par, cell["children"], bold=True, color="FFFFFF")
            else:
                add_inline(par, cell["children"], bold=False, color=None)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_diagram(node):
    # title
    for ch in node["children"]:
        if isinstance(ch, dict) and ch["tag"] == "div" and "dtitle" in ch["attrs"].get("class", ""):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, ch["children"], bold=True, color=NAVY_1, size=9)
            p.paragraph_format.space_after = Pt(4)
    for layer in [c for c in node["children"] if isinstance(c, dict) and c["tag"] == "div" and "layer" in c["attrs"].get("class", "")]:
        # layer label
        for ch in layer["children"]:
            if isinstance(ch, dict) and ch["tag"] == "div" and "lbl" in ch["attrs"].get("class", ""):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, ch["children"], bold=True, color=GRAY, size=7.5)
                p.paragraph_format.space_after = Pt(3)
        # rows of boxes + inter-layer arrows, in document order
        for child in layer["children"]:
            if not isinstance(child, dict) or child["tag"] != "div":
                continue
            cls = child["attrs"].get("class", "")
            if "row" in cls:
                items = [c for c in child["children"] if isinstance(c, dict) and c["tag"] == "div"
                         and ("box" in c["attrs"].get("class", "") or "arrow" in c["attrs"].get("class", ""))]
                if not items:
                    continue
                t = doc.add_table(rows=1, cols=len(items))
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for i, item in enumerate(items):
                    cell = t.cell(0, i)
                    cell.width = Cm(3.6)
                    icls = item["attrs"].get("class", "")
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after = Pt(0)
                    if "arrow" in icls:
                        add_inline(p, item["children"], bold=True, color=GOLD_BRIGHT)
                        continue
                    bg = None
                    for k, v in BOX_BG.items():
                        if k and k in icls:
                            bg = v
                            break
                    if bg is None:
                        bg = BOX_BG[None]
                    _shade_cell(cell, bg)
                    b_text = text_of([c for c in item["children"] if isinstance(c, dict) and c["tag"] == "b"])
                    small_text = text_of([c for c in item["children"] if isinstance(c, dict) and c["tag"] == "small"])
                    if b_text:
                        r = p.add_run(b_text)
                        _fmt(r, bold=True, color="FFFFFF", size=9)
                        p.add_run().add_break()
                    if small_text:
                        r = p.add_run(small_text)
                        _fmt(r, color="C9D1D9", size=7.5)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            elif "arrow" in cls:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, child["children"], bold=True, color=GOLD_BRIGHT)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_card(node):
    cls = node["attrs"].get("class", "")
    h5 = None
    ps = []
    for ch in node["children"]:
        if not isinstance(ch, dict):
            continue
        if ch["tag"] == "h5":
            h5 = ch
        elif ch["tag"] == "p":
            ps.append(ch)
    if h5:
        p = doc.add_paragraph()
        add_inline(p, h5["children"], bold=True, color=NAVY_1, size=10.5)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
    for body in ps:
        p = doc.add_paragraph()
        add_inline(p, body["children"], color=INK_SOFT, size=9.5)
        p.paragraph_format.space_after = Pt(5)


def render_stat(node):
    big = cap = None
    for ch in node["children"]:
        if not isinstance(ch, dict):
            continue
        if ch["tag"] == "span" and "big" in ch["attrs"].get("class", ""):
            big = ch
        elif ch["tag"] == "span" and "cap" in ch["attrs"].get("class", ""):
            cap = ch
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    if big:
        gold = "gold" in big["attrs"].get("class", "")
        add_inline(p, big["children"], bold=True, size=17,
                   color=GOLD if gold else NAVY_1)
    if cap:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(6)
        add_inline(p2, cap["children"], color=INK_SOFT, size=8.5)


def render_note(node):
    cls = node["attrs"].get("class", "")
    tone = next((t for t in ("blue", "gold", "green", "red") if t in cls), "blue")
    p = doc.add_paragraph()
    add_inline(p, node["children"])
    _shade_paragraph(p, NOTE_BG[tone])
    _border_left(p, NOTE_BORDER[tone])
    p.paragraph_format.space_after = Pt(8)


def render_cover(node):
    center = WD_ALIGN_PARAGRAPH.CENTER
    para([{"tag": "span", "attrs": {}, "children": ["PROJECT REPORT  ·  FINAL YEAR / CAPSTONE"]}],
         color=GRAY, size=9, space_after=2)
    para([{"tag": "span", "attrs": {}, "children": ["EKA — ENTERPRISE KNOWLEDGE ASSISTANT"]}],
         color=GRAY, size=9, space_after=26)
    para([{"tag": "span", "attrs": {}, "children": ["A RAG-POWERED DOCUMENT Q&A PLATFORM"]}],
         align=center, color=GOLD, size=10, space_after=18)
    p = doc.add_paragraph()
    p.alignment = center
    r = p.add_run("Enterprise Knowledge ")
    _fmt(r, bold=True, size=34, color=NAVY_1)
    r2 = p.add_run("Assistant")
    _fmt(r2, bold=True, size=34, color=GOLD)
    p.paragraph_format.space_after = Pt(14)
    para([{"tag": "span", "attrs": {}, "children": [
        "Upload company documents — policies, handbooks, budgets, runbooks — and ask questions "
        "in plain language. Get answers grounded in your documents, with clickable citations "
        "back to the exact source section."]}],
        align=center, color=INK_SOFT, size=11.5, space_after=16)
    para([{"tag": "span", "attrs": {}, "children": [
        "Upload  →  Index  →  Ask  →  Cited Answer"]}],
        align=center, bold=True, color=NAVY_1, size=12, space_after=30)
    para([{"tag": "span", "attrs": {}, "children": ["[Student Name / Team Members — confirm]"]}],
         align=center, bold=True, size=12.5, space_after=4)
    para([{"tag": "span", "attrs": {}, "children": [
        "[College / University Name — confirm]  ·  Department of Computer Science & Engineering"]}],
        align=center, color=INK_SOFT, size=10.5, space_after=4)
    para([{"tag": "span", "attrs": {}, "children": ["ACADEMIC SESSION 2025 – 2026"]}],
         align=center, color=GOLD, size=10, space_after=34)
    para([{"tag": "span", "attrs": {}, "children": ["Guided by: [Guide / Mentor Name — confirm]"]}],
         color=GRAY, size=9, space_after=2)
    para([{"tag": "span", "attrs": {}, "children": ["Tech: FastAPI · SQLite · FAISS · BM25 · MiniLM · Groq"]}],
         color=GRAY, size=9, space_after=0)
    page_break()


def add_toc_field():
    note = doc.add_paragraph()
    add_inline(note, [{"tag": "span", "attrs": {}, "children": [
        "In Word: right-click the contents below and choose “Update Field” (or press Ctrl+A then F9) "
        "to generate the page numbers."]}], italic=True, color=GRAY, size=9)
    p = doc.add_paragraph()
    run = p.add_run()
    r = run._r
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-3" \\h \\z \\u'
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "Table of contents — right-click and choose Update Field."
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for el in (f1, it, f2, t, f3):
        r.append(el)


def add_footer():
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("EKA — Enterprise Knowledge Assistant · Project Report · Page ")
    _fmt(r1, color=GRAY, size=8.5)
    r2 = fp.add_run()
    _fmt(r2, color=GRAY, size=8.5)
    rr = r2._r
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    for el in (f1, it, f2):
        rr.append(el)


def render_block(node):
    """Render one top-level HTML node (a .sheet div) into the docx."""
    tag = node["tag"]
    cls = node["attrs"].get("class", "")
    if tag == "script" or tag == "style" or tag in ("meta", "title", "link", "head"):
        return
    if tag != "div":
        return
    if cls and "cover" in cls:
        render_cover(node)
        return
    if "no-print" in cls or "topbar" in cls or "footer-note" in cls:
        return
    for ch in node["children"]:
        if isinstance(ch, str):
            continue
        render_item(ch)


_first_heading = [True]


def render_item(node):
    tag = node["tag"]
    cls = node["attrs"].get("class", "")

    if tag == "h2" and "chapter" in cls:
        if not _first_heading[0]:
            page_break()
        _first_heading[0] = False
        heading(1, node["children"])
        return
    if tag == "h3":
        heading(2, node["children"])
        return
    if tag == "h4":
        heading(3, node["children"])
        return
    if tag == "p":
        if "lead" in cls:
            para(node["children"], color=INK_SOFT, size=11)
        elif "codecap" in cls:
            para(node["children"], italic=True, color=INK_SOFT, size=8.5)
        elif "keywords" in cls:
            para(node["children"])
        else:
            para(node["children"])
        return
    if tag == "pre" and "code" in cls:
        code_block(node["children"])
        return
    if tag == "table":
        render_table(node)
        return
    if tag == "ul":
        for li in [c for c in node["children"] if isinstance(c, dict) and c["tag"] == "li"]:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, li["children"])
        return
    if tag == "ol":
        for li in [c for c in node["children"] if isinstance(c, dict) and c["tag"] == "li"]:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, li["children"])
        return
    if tag == "div":
        if "diagram" in cls:
            render_diagram(node)
            return
        if "card" in cls:
            render_card(node)
            return
        if "stat" in cls:
            render_stat(node)
            return
        if "note" in cls:
            render_note(node)
            return
        if "toc" in cls:
            add_toc_field()
            return
        if "flowline" in cls:
            parts = []
            for c in node["children"]:
                if not isinstance(c, dict):
                    continue
                if "arrow" in c["attrs"].get("class", ""):
                    parts.append(("arrow", " → "))
                elif "box" in c["attrs"].get("class", ""):
                    name = text_of([x for x in c["children"] if isinstance(x, dict) and x["tag"] == "b"])
                    small = text_of([x for x in c["children"] if isinstance(x, dict) and x["tag"] == "small"])
                    parts.append(("box", f"{name} ({small})" if small else name))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for kind, txt in parts:
                r = p.add_run(txt)
                if kind == "arrow":
                    _fmt(r, bold=True, color=GOLD_BRIGHT)
                else:
                    _fmt(r, bold=True, color=NAVY_1)
            p.paragraph_format.space_after = Pt(10)
            return
        # generic container: recurse
        for ch in node["children"]:
            if isinstance(ch, dict):
                render_item(ch)
        return


# ---------------------------------------------------------------- build
def main() -> None:
    html = HTML_SRC.read_text(encoding="utf-8")
    tb = TreeBuilder()
    tb.feed(html)

    # root > html > body
    html_node = next((c for c in tb.root["children"] if isinstance(c, dict) and c["tag"] == "html"), tb.root)
    body_node = next((c for c in html_node["children"] if isinstance(c, dict) and c["tag"] == "body"), html_node)
    sheets = [c for c in body_node["children"] if isinstance(c, dict) and c["tag"] == "div"
              and "sheet" in c["attrs"].get("class", "")]
    for sheet in sheets:
        render_block(sheet)

    add_footer()
    doc.save(DOCX_OUT)
    print(f"Saved {DOCX_OUT} ({DOCX_OUT.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()